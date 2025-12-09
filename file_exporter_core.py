"""
File Exporter Core Logic
========================
Core functions for scanning directories and extracting file metadata.
This module contains no GUI code and can be imported independently.

Requirements:
    - pandas (pip install pandas)
    - openpyxl (pip install openpyxl)
    - python-docx (pip install python-docx) - optional, for Word doc authors
    - python-pptx (pip install python-pptx) - optional, for PowerPoint authors
"""

import os
import datetime
import pandas as pd

# ============================================================
# OPTIONAL IMPORTS FOR FILE AUTHOR EXTRACTION
# These libraries allow reading author metadata from Office files
# The program will still work if they're not installed
# ============================================================
try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


def get_file_dates(filepath):
    """
    Get the creation and modification dates of a file.
    
    Args:
        filepath: Full path to the file
        
    Returns:
        tuple: (creation_date, modified_date) as formatted strings
    """
    try:
        stat = os.stat(filepath)
        
        # Get modification time (works on all platforms)
        modified_date = datetime.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        
        # Get creation time
        # On Windows: st_ctime is creation time
        # On Unix: st_ctime is last metadata change, st_birthtime is creation (if available)
        if hasattr(stat, 'st_birthtime'):
            # macOS
            creation_date = datetime.datetime.fromtimestamp(stat.st_birthtime).strftime("%Y-%m-%d %H:%M:%S")
        else:
            # Windows uses st_ctime as creation time
            creation_date = datetime.datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
        
        return creation_date, modified_date
    except Exception:
        return None, None


def get_file_author(filepath):
    """
    Attempt to extract the author metadata from a file.
    
    Supports:
        - Excel files (.xlsx, .xlsm)
        - Word documents (.docx)
        - PowerPoint files (.pptx)
    
    Args:
        filepath: Full path to the file
        
    Returns:
        str: Author name or None if not available
    """
    ext = os.path.splitext(filepath)[1].lower()
    
    try:
        # ----------------------------------------------------
        # EXCEL FILES (.xlsx, .xlsm)
        # Uses openpyxl to read document properties
        # ----------------------------------------------------
        if ext in ['.xlsx', '.xlsm'] and HAS_OPENPYXL:
            wb = load_workbook(filepath, read_only=True, data_only=True)
            author = wb.properties.creator
            wb.close()
            return author
        
        # ----------------------------------------------------
        # WORD DOCUMENTS (.docx)
        # Uses python-docx to read core properties
        # ----------------------------------------------------
        elif ext == '.docx' and HAS_DOCX:
            doc = Document(filepath)
            return doc.core_properties.author
        
        # ----------------------------------------------------
        # POWERPOINT FILES (.pptx)
        # Uses python-pptx to read core properties
        # ----------------------------------------------------
        elif ext == '.pptx' and HAS_PPTX:
            prs = Presentation(filepath)
            return prs.core_properties.author
        
        # ----------------------------------------------------
        # UNSUPPORTED FILE TYPES
        # Return None for files we can't extract author from
        # ----------------------------------------------------
        else:
            return None
            
    except Exception:
        # If anything goes wrong, just return None
        return None


def parse_extensions(extension_string):
    """
    Parse a comma-separated string of file extensions.
    
    Args:
        extension_string: Comma-separated extensions (e.g., ".psd, tif, .pdf")
        
    Returns:
        list: List of lowercase extensions with dots (e.g., [".psd", ".tif", ".pdf"])
    """
    if not extension_string or not extension_string.strip():
        return []
    
    extensions = [
        e.strip().lower() if e.strip().startswith(".") else f".{e.strip().lower()}" 
        for e in extension_string.split(",")
    ]
    return extensions


def parse_folder_structure(full_path, base_path, title_case=True):
    """
    Parse a full path into individual folder names relative to a base path.
    
    Args:
        full_path: Full directory path
        base_path: Base path to strip from the beginning
        title_case: Whether to convert folder names to title case
        
    Returns:
        list: List of folder names
    """
    relative_path = full_path.replace(base_path, "").lstrip("\\").lstrip("/")
    folders = [f for f in relative_path.split(os.sep) if f]
    
    if title_case:
        folders = [f.title() for f in folders]
    
    return folders


def get_root_folder_name(directory):
    """
    Get the name of the top-level folder from a directory path.
    
    Args:
        directory: Full directory path
        
    Returns:
        str: Name of the top-level folder
    """
    return os.path.basename(os.path.normpath(directory))


def build_row(filename, dirpath, root_name, folders, folder_cols, 
              include_dates=True, include_author=True):
    """
    Build a data row dictionary for a single file.
    
    Args:
        filename: Name of the file
        dirpath: Directory path containing the file
        root_name: Label for the root folder
        folders: List of parsed folder names
        folder_cols: Number of folder columns to include
        include_dates: Whether to include date columns
        include_author: Whether to include author column
        
    Returns:
        dict: Row data dictionary
    """
    filepath = os.path.join(dirpath, filename)
    
    row = {
        "RootFolder": root_name,
        "FolderString": dirpath,
        "FullPath": filepath,          # Full path including filename
        "FileName": filename,
    }
    
    # Add dynamic folder columns
    for i in range(folder_cols):
        row[f"Folder{i + 1}"] = folders[i] if i < len(folders) else None
    
    # Add date columns
    if include_dates:
        creation_date, modified_date = get_file_dates(filepath)
        row["DateCreated"] = creation_date
        row["DateModified"] = modified_date
    
    # Add author column
    if include_author:
        row["Author"] = get_file_author(filepath)
    
    return row


def scan_directory(directory, root_name=None, folder_cols=3, title_case=True,
                   extensions=None, include_dates=True, include_author=True,
                   progress_callback=None, cancel_check=None):
    """
    Scan a directory and collect file information.
    
    Args:
        directory: Path to directory to scan
        root_name: Label for root folder (defaults to top-level folder name)
        folder_cols: Number of folder columns to include
        title_case: Whether to convert folder names to title case
        extensions: List of file extensions to filter (empty = all files)
        include_dates: Whether to include date columns
        include_author: Whether to include author column
        progress_callback: Function to call with progress updates (receives file count)
        cancel_check: Function that returns True if scan should be cancelled
        
    Returns:
        list: List of row dictionaries
    """
    # Default root name to the top-level folder name
    if root_name is None:
        root_name = get_root_folder_name(directory)
    
    if extensions is None:
        extensions = []
    
    files = []
    file_count = 0
    
    for dirpath, dirs, filenames in os.walk(directory):
        # Check for cancellation
        if cancel_check and cancel_check():
            break
        
        for filename in filenames:
            # Check for cancellation
            if cancel_check and cancel_check():
                break
            
            # Extension filtering
            if extensions:
                ext = os.path.splitext(filename)[1].lower()
                if ext not in extensions:
                    continue
            
            # Parse folder structure
            folders = parse_folder_structure(dirpath, directory, title_case)
            
            # Build row
            row = build_row(
                filename=filename,
                dirpath=dirpath,
                root_name=root_name,
                folders=folders,
                folder_cols=folder_cols,
                include_dates=include_dates,
                include_author=include_author
            )
            
            files.append(row)
            file_count += 1
            
            # Progress callback
            if progress_callback and file_count % 100 == 0:
                progress_callback(file_count)
    
    return files


def export_to_excel(files, output_file):
    """
    Export file data to an Excel file.
    
    Args:
        files: List of row dictionaries
        output_file: Path to output Excel file
        
    Returns:
        int: Number of files exported
    """
    df = pd.DataFrame(files)
    df.to_excel(output_file, index=False)
    return len(df)
